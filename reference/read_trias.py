"""Read yearly results from the recalculated trias workbook; entry point `read_trias`."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

FIRST_COL, N_YEARS = 6, 35  # column F = year 1 ... AN = year 35

# Signs as in the workbook: costs and repayments are negative.
ROWS = {
    "Finanzplan": {
        "rent_gross": 8,
        "maintenance": 9,
        "vacancy": 10,
        "legal_admin": 52,
        "direct_loan_interest": 104,
        "bank_interest": 106,
        "subsidy_interest": 108,
        "liquidity": 122,
    },
    "Kapital & Finanzierung": {
        "bank_repayment": 79,
        "bank_balance": 81,
        "subsidy_repayment": 121,
        "subsidy_balance": 123,
    },
}


def read_trias(path: Path) -> pd.DataFrame:
    """Return one column per quantity in ROWS, indexed by operating year 1..35."""
    wb = load_workbook(path, data_only=True, read_only=True)
    data = {}
    for sheet, rows in ROWS.items():
        ws = wb[sheet]
        for name, row in rows.items():
            cells = next(
                ws.iter_rows(
                    min_row=row,
                    max_row=row,
                    min_col=FIRST_COL,
                    max_col=FIRST_COL + N_YEARS - 1,
                    values_only=True,
                )
            )
            data[name] = [float(v or 0.0) for v in cells]
    wb.close()
    return pd.DataFrame(data, index=pd.RangeIndex(1, N_YEARS + 1, name="t"))


if __name__ == "__main__":
    print(read_trias(Path(__file__).with_name("trias_reference_case.xlsx")).round(0).to_string())
